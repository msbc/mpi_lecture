"""PptxEngine: builds a PowerPoint deck with python-pptx, styled like the Keynote deck.

The slide is 1920 x 1080 pt, so one Keynote pixel is one point and Keynote font sizes carry over unchanged.
Everything is placed at the same coordinates as in Keynote. Titles are the slides' title placeholders, so
PowerPoint's outline and navigation work; the rest are plain shapes.
"""

import os
import shutil
import subprocess

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

from .engine import BuildEngine
from .markup import inline_code
from .style import (
    BODY_SIZE,
    CARD_BORDER,
    CHAR,
    CODE,
    COVER_FOOTER,
    COVER_SUB,
    COVER_TITLE,
    CW,
    HEAD,
    HEAD_TITLE,
    LINE,
    MONO,
    MONO_BOLD,
    SANS,
    SANS_REGULAR,
    SECTION_TITLE,
    STATEMENT,
    SUBTITLE,
    TEXT,
    TITLE,
    X0,
    Y0,
    YMAX,
    H,
    W,
)

# Helvetica Neue stands in for the theme's Graphik and Produkt, which only Macs have; PowerPoint on a
# machine without these fonts substitutes its own
FONTS = {
    SANS: "Helvetica Neue Light",
    SANS_REGULAR: "Helvetica Neue",
    HEAD: "Helvetica Neue Light",
    HEAD_TITLE: "Helvetica Neue Light",
    MONO: "Menlo",
    MONO_BOLD: "Menlo",
}
LAYOUT_TITLE_ONLY = 5  # of the default template
PARA_SKIP = 0.5  # space after a paragraph of a text item, in em
ITEM_SKIP = 0.55  # space between the bullets of a body, in em
TEXT_LINE = 1.1  # line spacing of proportional text, as a multiple of single spacing
AVG_CHAR = 0.5  # average advance of the sans fonts / font size, to estimate how much text fits
EQ_DPI = 600  # equations are PDFs, which PowerPoint can't place; rasterize them at this resolution


def emu(px):
    return Emu(round(px * 12700))


def rgb(h):
    return RGBColor.from_string(h.lstrip("#"))


def fit_size(paragraphs, w, h, size, line=TEXT_LINE, skip=ITEM_SKIP, indent=0.0):
    """Largest font size <= size, shrunk by 5% steps, at which the paragraphs (wrapped at an estimated
    character width) fit the box w x h, which PowerPoint would otherwise only do once the text is edited."""
    while size > 8:
        usable = w - indent * size
        lines = sum(max(1, -(-int(len(p) * AVG_CHAR * size) // int(usable))) for p in paragraphs)
        if lines * line * size + skip * size * (len(paragraphs) - 1) <= h:
            break
        size *= 0.95
    return size


def rasterize(path):
    """A PDF's PNG (made next to it, cached); other files are returned unchanged."""
    if not path.endswith(".pdf"):
        return path
    png = path[:-4] + f"_{EQ_DPI}.png"
    if not os.path.exists(png):
        if not shutil.which("pdftoppm"):
            raise SystemExit(
                "The PowerPoint engine needs pdftoppm (poppler) to place the equations: brew install poppler"
            )
        subprocess.run(["pdftoppm", "-png", "-r", str(EQ_DPI), "-singlefile", path, png[:-4]], check=True)
    return png


class PptxEngine(BuildEngine):
    default_output = "Intro_MPI.pptx"
    protect_existing = False  # a generated deck has no hand edits to lose
    fonts = FONTS

    def __init__(self, output):
        super().__init__(output)
        self.prs = None
        self.slide = None
        self.master = None

    # ------------------------------------------------------------- document
    def begin(self):
        self.prs = Presentation()
        self.prs.slide_width, self.prs.slide_height = emu(W), emu(H)

    def finish(self):
        self.end_slide()

    def source(self):
        raise NotImplementedError("a .pptx is built directly, with no script to print")

    def render(self):
        self.prs.save(self.output)
        print(self.output)
        return 0

    # ------------------------------------------------------------ text boxes
    @staticmethod
    def place(shape, x, y, w, h):
        shape.left, shape.top, shape.width, shape.height = emu(x), emu(y), emu(w), emu(h)

    @staticmethod
    def frame(shape, valign="top", wrap=True):
        """Text frame with no margins and no autofit."""
        tf = shape.text_frame
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        tf.word_wrap = wrap
        tf.auto_size = MSO_AUTO_SIZE.NONE
        tf.vertical_anchor = {"top": MSO_ANCHOR.TOP, "middle": MSO_ANCHOR.MIDDLE, "bottom": MSO_ANCHOR.BOTTOM}[valign]
        return tf

    def run(self, paragraph, text, font, size, color, bold=False, link=None):
        r = paragraph.add_run()
        r.text = text
        r.font.name = self.fonts[font]
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.color.rgb = rgb(color)
        if link:
            r.hyperlink.address = link
            r.font.underline = True
        return r

    def inline_runs(self, paragraph, s, font, size, color, link=None):
        """Runs for s, with `backtick` text in Menlo."""
        plain, ranges = inline_code(s)
        pos = 0
        for a, b in ranges:
            if plain[pos : a - 1]:
                self.run(paragraph, plain[pos : a - 1], font, size, color, link=link)
            self.run(paragraph, plain[a - 1 : b], MONO, size, color, link=link)
            pos = b
        if plain[pos:]:
            self.run(paragraph, plain[pos:], font, size, color, link=link)

    def fill(self, shape, s, font, size, color, valign="top", align="left", link=None):
        """Text s (paragraphs split at newlines) in a shape's frame."""
        tf = self.frame(shape, valign)
        for i, para in enumerate(s.split("\n")):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment = {"left": PP_ALIGN.LEFT, "right": PP_ALIGN.RIGHT, "center": PP_ALIGN.CENTER}[align]
            p.line_spacing = TEXT_LINE
            p.space_after = Pt(PARA_SKIP * size)
            self.inline_runs(p, para, font, size, color, link)

    def box_text(self, box, s, shape=None):
        """Text s in a master's text box; shrunk to fit the width if the box is a one-line box."""
        size = box.size
        if box.shrink:
            plain, ranges = inline_code(s)
            code = sum(b - a + 1 for a, b in ranges)  # Menlo is wider than the sans font
            em = AVG_CHAR * (len(plain) - code) + CHAR * code
            size = min(size, box.w / em)
        if shape is None:
            shape = self.slide.shapes.add_textbox(0, 0, 0, 0)
        self.place(shape, box.x, box.y, box.w, box.h)
        self.fill(shape, s, box.font, size, TEXT, box.valign, box.halign)

    # ------------------------------------------------------------ primitives
    def end_slide(self):
        pass

    def start_slide(self, master):
        self.master = master
        self.slide = self.prs.slides.add_slide(self.prs.slide_layouts[LAYOUT_TITLE_ONLY])

    def title(self, title, sub=None):
        if self.master == "Section":
            self.box_text(SECTION_TITLE, title, self.slide.shapes.title)
            return
        self.box_text(TITLE, title, self.slide.shapes.title)
        if sub is not None:
            self.box_text(SUBTITLE, sub)

    def title_slide(self, title, sub, footer):
        self.box_text(COVER_TITLE, title, self.slide.shapes.title)
        self.box_text(COVER_SUB, sub)
        self.box_text(COVER_FOOTER, footer)

    def statement(self, text):
        self.box_text(STATEMENT, text, self.slide.shapes.title)

    def body(self, bullets, x=X0, y=Y0, w=CW, h=None, size=None):
        h = h or (YMAX - y)
        size = fit_size([inline_code(b)[0] for b in bullets], w, h, size or BODY_SIZE, skip=ITEM_SKIP, indent=1.1)
        shape = self.slide.shapes.add_textbox(0, 0, 0, 0)
        self.place(shape, x, y, w, h)
        tf = self.frame(shape)
        for i, b in enumerate(bullets):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.line_spacing = TEXT_LINE
            p.space_after = Pt(ITEM_SKIP * size)
            pPr = p._p.get_or_add_pPr()
            pPr.set("marL", str(round(1.1 * size * 12700)))
            pPr.set("indent", str(round(-1.1 * size * 12700)))
            bu = etree.SubElement(pPr, qn("a:buFont"))
            bu.set("typeface", "Arial")
            etree.SubElement(pPr, qn("a:buChar")).set("char", "•")
            self.inline_runs(p, b, SANS, size, TEXT)

    def text(self, s, x, y, w, size=36, font=SANS, color=TEXT, align="left", link=None):
        h = size * TEXT_LINE * (s.count("\n") + 1) * 1.4  # room for wrapped lines; the box has no fill
        shape = self.slide.shapes.add_textbox(0, 0, 0, 0)
        self.place(shape, x, y, w, h)
        self.fill(shape, s, font, size, color, align=align, link=link)

    def code_text(self, code, x, y, w, size, spans):
        colors, bold = [CODE] * len(code), [False] * len(code)
        for a, b, col, bf in spans:  # later spans win, as in Keynote
            for i in range(a - 1, b):
                colors[i] = col
                bold[i] = bold[i] or bf
        shape = self.slide.shapes.add_textbox(0, 0, 0, 0)
        lines = code.split("\n")
        self.place(shape, x, y, w, LINE * size * len(lines))
        tf = self.frame(shape, wrap=False)
        pos = 0
        for i, line in enumerate(lines):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.line_spacing = Pt(LINE * size)
            p.space_after = 0
            col, runs = 0, []  # runs of equal color and weight: [text, color, bold]
            for j, ch in enumerate(line):
                k = (colors[pos + j], bold[pos + j])
                if ch == "\t":
                    ch, n = " ", 8 - col % 8
                else:
                    n = 1
                if runs and (runs[-1][1], runs[-1][2]) == k:
                    runs[-1][0] += ch * n
                else:
                    runs.append([ch * n, *k])
                col += n
            for text, c, bf in runs:
                self.run(p, text, MONO_BOLD if bf else MONO, size, c, bold=bf)
            pos += len(line) + 1

    def card(self, x, y, w, h, alpha=0.82):
        shape = self.slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, emu(x), emu(y), emu(w), emu(h))
        shape.adjustments[0] = min(0.5, 22 / min(w, h))
        # the theme style would add a gradient fill and a shadow in some renderers; everything is set explicitly
        shape._element.remove(shape._element.find(qn("p:style")))
        shape.fill.solid()
        shape.fill.fore_color.rgb = RGBColor(255, 255, 255)
        shape.line.color.rgb = rgb(CARD_BORDER)
        shape.line.width = Pt(2)
        # python-pptx has no API for transparency: add an alpha to the fill's color
        clr = shape._element.spPr.find(qn("a:solidFill"))[0]
        etree.SubElement(clr, qn("a:alpha")).set("val", str(round(alpha * 100000)))

    def image(self, path, x, y, w, h):
        self.slide.shapes.add_picture(rasterize(path), emu(x), emu(y), emu(w), emu(h))

    def notes(self, s):
        self.slide.notes_slide.notes_text_frame.text = s
